from __future__ import annotations

import base64
import os
from io import BytesIO
from pathlib import Path

from cryptography.fernet import Fernet
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = BASE_DIR / "word" / "base_template.docx"
ENCRYPTED_TEMPLATE_PATH = BASE_DIR / "word" / "base_template.docx.enc"
WESTERN_FONT = "Times New Roman"


def _open_template_document():
    """Load the private Word template from disk or decrypt its deployment copy in memory."""
    if TEMPLATE_PATH.exists():
        return Document(str(TEMPLATE_PATH))
    key = os.getenv("TEMPLATE_KEY", "").strip()
    if not key or not ENCRYPTED_TEMPLATE_PATH.exists():
        raise RuntimeError("Word template is not configured")
    decrypted = Fernet(key.encode("ascii")).decrypt(ENCRYPTED_TEMPLATE_PATH.read_bytes())
    return Document(BytesIO(decrypted))


def _fmt_num(value, empty="//"):
    return empty if value is None or value == "" else str(value)


def _force_western_font(run):
    """Make dynamic Vietnamese/Latin text match the Word template's Times New Roman."""
    run.font.name = WESTERN_FONT
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), WESTERN_FONT)
    r_fonts.set(qn("w:hAnsi"), WESTERN_FONT)
    r_fonts.set(qn("w:cs"), WESTERN_FONT)


def _set_run_value(cell, value: str, run_index: int = 0, western: bool = True):
    p = cell.paragraphs[0]
    while len(p.runs) <= run_index:
        p.add_run()
    p.runs[run_index].text = value
    if western:
        _force_western_font(p.runs[run_index])
    for i, run in enumerate(p.runs):
        if i != run_index:
            run.text = ""


def _set_bilingual_inline_cell(cell, zh: str, vi: str, blank="//"):
    """Write Chinese + one literal space + Vietnamese using separate formatted runs."""
    zh = zh or blank
    vi = vi or blank
    p = cell.paragraphs[0]
    if not p.runs:
        p.add_run()
    chinese_run = p.runs[0]
    chinese_run.text = f"{zh} "

    if len(p.runs) >= 2:
        vietnamese_run = p.runs[1]
    else:
        vietnamese_run = p.add_run()
    vietnamese_run.text = vi
    _force_western_font(vietnamese_run)

    # Match the original sample's emphasis for bilingual values.
    if chinese_run.bold is not False:
        chinese_run.bold = True
    vietnamese_run.bold = True

    for run in p.runs[2:]:
        run.text = ""


def _set_bilingual_cell(cell, zh: str, vi: str, blank="//"):
    zh = zh or blank
    vi = vi or blank
    while len(cell.paragraphs) < 2:
        cell.add_paragraph()
    for idx, value in enumerate((zh, vi)):
        p = cell.paragraphs[idx]
        if p.runs:
            p.runs[0].text = value
            for run in p.runs[1:]:
                run.text = ""
        else:
            p.add_run(value)
        if idx == 1:
            _force_western_font(p.runs[0])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for extra in cell.paragraphs[2:]:
        for run in extra.runs:
            run.text = ""


def _yes_no_cells(yes_cell, no_cell, value: bool):
    for cell, checked in ((yes_cell, value), (no_cell, not value)):
        p = cell.paragraphs[0]
        if p.runs:
            p.runs[0].text = "■" if checked else "□"


def _replace_text_in_paragraph(paragraph, old: str, new: str):
    full = "".join(r.text for r in paragraph.runs)
    if old not in full:
        return False
    pos = full.find(old)
    cursor = 0
    start_idx = 0
    end_idx = 0
    start_off = 0
    end_off = 0
    for i, run in enumerate(paragraph.runs):
        nxt = cursor + len(run.text)
        if cursor <= pos < nxt or (pos == nxt and i == len(paragraph.runs) - 1):
            start_idx = i
            start_off = pos - cursor
        end_pos = pos + len(old)
        if cursor < end_pos <= nxt:
            end_idx = i
            end_off = end_pos - cursor
            break
        cursor = nxt
    before = paragraph.runs[start_idx].text[:start_off]
    after = paragraph.runs[end_idx].text[end_off:]
    paragraph.runs[start_idx].text = before + new + after
    _force_western_font(paragraph.runs[start_idx])
    for i in range(start_idx + 1, end_idx + 1):
        paragraph.runs[i].text = ""
    return True


def _set_marital_status(cell, zh: str, vi: str):
    """Preserve the original label runs and keep Chinese/Vietnamese in their native formatting."""
    p = cell.paragraphs[0]
    zh = zh or "//"
    vi = vi or "//"
    if len(p.runs) >= 11:
        p.runs[7].text = zh
        p.runs[8].text = " "
        p.runs[9].text = ""
        p.runs[10].text = vi
        _force_western_font(p.runs[10])
        for run in p.runs[11:]:
            run.text = ""
        return
    p.add_run(f"{zh} ")
    vi_run = p.add_run(vi)
    _force_western_font(vi_run)


def _set_document_checkboxes(cell, passport: bool, judicial: bool, health: bool):
    p = cell.paragraphs[0]
    for run_idx, value in ((0, passport), (6, judicial), (9, health)):
        if run_idx < len(p.runs):
            p.runs[run_idx].text = "■" if value else "□"


def _remove_all_drawings(cell):
    for p in cell.paragraphs:
        for drawing in p._p.xpath(".//w:drawing"):
            parent = drawing.getparent()
            parent.remove(drawing)


def _photo_stream(photo_value: str | None):
    if not photo_value:
        return None
    if photo_value.startswith("data:") and ";base64," in photo_value:
        try:
            encoded = photo_value.split(";base64,", 1)[1]
            return BytesIO(base64.b64decode(encoded))
        except Exception:
            return None
    path = Path(photo_value)
    if path.exists() and path.is_file():
        return str(path)
    return None


def _set_photo(cell, photo_value: str | None):
    _remove_all_drawings(cell)
    source = _photo_stream(photo_value)
    if source is None:
        return
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(source, width=Cm(3.6), height=Cm(4.8))


def export_profile(profile, output_path: str | Path):
    doc = _open_template_document()
    table = doc.tables[0]

    # Reinforce the western font used by the original template.
    normal = doc.styles["Normal"]
    normal.font.name = WESTERN_FONT
    normal_rpr = normal.element.get_or_add_rPr()
    normal_fonts = normal_rpr.get_or_add_rFonts()
    normal_fonts.set(qn("w:ascii"), WESTERN_FONT)
    normal_fonts.set(qn("w:hAnsi"), WESTERN_FONT)
    normal_fonts.set(qn("w:cs"), WESTERN_FONT)

    # Header code.
    p = table.cell(1, 0).paragraphs[0]
    if len(p.runs) > 6:
        p.runs[6].text = profile.code or "MS"
        _force_western_font(p.runs[6])

    _set_run_value(table.cell(2, 1), profile.full_name_vi or "//")
    _set_run_value(table.cell(2, 9), profile.full_name_zh or "//", western=False)
    _set_run_value(table.cell(3, 1), profile.birth_date or "//")
    _set_run_value(table.cell(3, 13), _fmt_num(profile.age))
    _set_run_value(
        table.cell(4, 1),
        f"{_fmt_num(profile.height_cm)}CM" if profile.height_cm is not None else "//",
    )
    _set_run_value(
        table.cell(4, 13),
        f"{_fmt_num(profile.weight_kg)} KG" if profile.weight_kg is not None else "//",
    )

    _set_bilingual_inline_cell(table.cell(5, 1), profile.education_zh, profile.education_vi)

    # Religion and foreign-language rows are intentionally copied unchanged from base_template.docx.
    # The public form no longer asks users to enter these fixed values.
    _set_marital_status(table.cell(6, 5), profile.marital_status_zh, profile.marital_status_vi)

    _set_run_value(table.cell(7, 1), profile.province_zh or "//", western=False)
    _set_run_value(table.cell(7, 9), profile.province_vi or "//")

    _replace_fixed_tail(table.cell(9, 0), 6, _fmt_num(profile.children_count))
    _replace_fixed_tail(table.cell(9, 9), 4, _fmt_num(profile.sons_count))
    _replace_fixed_tail(table.cell(9, 14), 5, _fmt_num(profile.daughters_count))

    _set_run_value(table.cell(11, 3), profile.father_name or "//")
    _set_run_value(table.cell(11, 12), _fmt_num(profile.father_birth_year))
    _set_bilingual_inline_cell(table.cell(11, 15), profile.father_job_zh, profile.father_job_vi)

    _set_run_value(table.cell(12, 3), profile.mother_name or "//")
    _set_run_value(table.cell(12, 12), _fmt_num(profile.mother_birth_year))
    _set_bilingual_inline_cell(table.cell(12, 15), profile.mother_job_zh, profile.mother_job_vi)

    _set_run_value(table.cell(13, 3), profile.spouse_name or "//")
    _set_run_value(table.cell(13, 12), _fmt_num(profile.spouse_birth_year))
    if profile.spouse_job_vi:
        _set_bilingual_inline_cell(table.cell(13, 15), profile.spouse_job_zh, profile.spouse_job_vi)
    else:
        _set_run_value(table.cell(13, 15), "//")

    _set_run_value(table.cell(14, 4), _fmt_num(profile.siblings_count))
    _set_run_value(table.cell(14, 15), _fmt_num(profile.birth_order))

    rel = table.cell(15, 9).paragraphs[0]
    if len(rel.runs) > 4:
        rel.runs[0].text = "■" if profile.relatives_in_taiwan else "□"
        rel.runs[4].text = "□" if profile.relatives_in_taiwan else "■"

    _set_bilingual_cell(table.cell(17, 2), profile.work_country_zh, profile.work_country_vi)
    _set_run_value(table.cell(17, 7), profile.work_period or "//")
    _set_bilingual_cell(table.cell(17, 11), profile.work_description_zh, profile.work_description_vi)

    for cell in (table.cell(18, 2), table.cell(18, 7), table.cell(18, 11)):
        for para in cell.paragraphs:
            for run in para.runs:
                run.text = ""

    tw = table.cell(19, 11).paragraphs[0]
    if len(tw.runs) > 4:
        tw.runs[0].text = "■" if profile.taiwan_work_experience else "□"
        tw.runs[4].text = "□" if profile.taiwan_work_experience else "■"

    _set_document_checkboxes(
        table.cell(20, 3),
        profile.has_passport,
        profile.has_judicial_record,
        profile.has_health_exam,
    )

    _yes_no_cells(table.cell(21, 3), table.cell(21, 8), profile.smoking)
    _yes_no_cells(table.cell(22, 3), table.cell(22, 8), profile.alcohol)
    _yes_no_cells(table.cell(23, 3), table.cell(23, 8), profile.color_blind)
    _yes_no_cells(table.cell(24, 3), table.cell(24, 8), profile.tattoos)
    _yes_no_cells(table.cell(25, 3), table.cell(25, 8), profile.habit)

    date_cell = table.cell(26, 0)
    _replace_text_in_paragraph(
        date_cell.paragraphs[0], "{{FORM_DATE}}", profile.form_date or "//"
    )
    sig_cell = table.cell(26, 10)
    _replace_text_in_paragraph(
        sig_cell.paragraphs[0], "{{FULL_NAME}}", profile.full_name_vi or "//"
    )

    _set_photo(table.cell(2, 15), profile.photo_path)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def _replace_fixed_tail(cell, start_run: int, value: str):
    """Replace only the dynamic tail of a fixed bilingual label and keep Latin values in TNR."""
    p = cell.paragraphs[0]
    if start_run >= len(p.runs):
        run = p.add_run(value)
        _force_western_font(run)
        return
    p.runs[start_run].text = value
    _force_western_font(p.runs[start_run])
    for run in p.runs[start_run + 1 :]:
        run.text = ""
